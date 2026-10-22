"""Editable PowerPoint with the method diagrams as native shapes (architecture, pipeline, tasks) plus every figure as a
picture slide.    .venv/bin/python docs/paper/make_pptx.py  ->  docs/paper/flapsim_paper_figures.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parents[2]
FIGS = ROOT/'docs/paper/figs'
OUT = ROOT/'docs/paper/flapsim_paper_figures.pptx'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
C = dict(learn=RGBColor(0x1f, 0x8a, 0x4c), gate=RGBColor(0xc9, 0xa2, 0x27), muscle=RGBColor(0xd4, 0x78, 0x1f), reflex=RGBColor(0x2f, 0x6d, 0xb5),
         classic=RGBColor(0x66, 0x66, 0x66), blind=RGBColor(0x7b, 0x4f, 0xa8), robust=RGBColor(0xc2, 0x3b, 0x3b), ink=RGBColor(0x22, 0x22, 0x22), white=RGBColor(0xff, 0xff, 0xff))


def box(sl, x, y, w, h, title, body='', color=C['classic'], fill=None, fs=11):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = .08
    s.fill.solid()
    s.fill.fore_color.rgb = fill or RGBColor(0xff, 0xff, 0xff)
    s.line.color.rgb = color
    s.line.width = Pt(1.5)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.bold = True
    r.font.size = Pt(fs)
    r.font.color.rgb = color
    if body:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = body
        r2.font.size = Pt(fs-2)
        r2.font.color.rgb = C['ink']
    return s


def arrow(sl, a, b, color=None, text=''):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    c.begin_connect(a, 3)
    c.end_connect(b, 1)
    c.line.color.rgb = color or C['ink']
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd', {'type': 'triangle'})
    ln.append(tail)
    if text:
        tb = sl.shapes.add_textbox(Emu((a.left+a.width+b.left)//2-Inches(.6)), Emu(min(a.top, b.top)-Inches(.3)), Inches(1.2), Inches(.3))
        tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(9)
        r.font.color.rgb = C['ink']
    return c


def title(sl, t, sub=''):
    tb = sl.shapes.add_textbox(Inches(.4), Inches(.2), Inches(12.5), Inches(.6))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = t
    r.font.size = Pt(20)
    r.font.bold = True
    if sub:
        p = tb.text_frame.add_paragraph()
        r2 = p.add_run()
        r2.text = sub
        r2.font.size = Pt(11)
        r2.font.color.rgb = C['classic']


# ---- slide 1: architecture (one 2 ms tick)
sl = prs.slides.add_slide(BLANK)
title(sl, 'Three-layer controller (one 2 ms tick)', 'Editable shapes; colours: green = learned descending layer, yellow = command gate, orange = steering muscle, blue = reflex PID, grey = classical')
imu = box(sl, .4, 2.2, 1.9, 1.2, 'IMU + ToF', '11-D reading, 500 Hz, noisy')
cmd = box(sl, .4, 4.4, 1.9, 1.3, 'Task command', 'evade channel or dash/redirect channel; unit vector along φ')
feat = box(sl, 2.8, 1.5, 1.7, 1.5, 'Features', '25-D: 16 rectified IMU, bias, 8 command channels', C['learn'])
dl = box(sl, 4.7, 1.5, 1.5, 1.5, 'Delay lines', '11 taps, 0–20 ms', C['learn'])
hid = box(sl, 6.4, 1.2, 1.9, 1.0, 'LIF hidden layer', '8 neurons, 2 sub-steps', C['learn'])
outp = box(sl, 6.4, 2.35, 1.9, 1.3, 'Outputs', 'via 11 axonal delays: 2 spiking → pulse; 2 graded g_r, g_p', C['learn'])
gate = box(sl, 8.7, 1.4, 2.0, 2.0, 'Command gate', 'learned pulse and set-points in charge while the command is on + 1 s tail; otherwise classical base law, set-points 0', C['gate'], RGBColor(0xff, 0xf8, 0xe1))
refr = box(sl, 11.0, 1.4, 1.9, 1.0, 'Refractory 37 ms', '≤ 1 learned onset per wingbeat', C['gate'], RGBColor(0xff, 0xf8, 0xe1))
musc = box(sl, 11.0, 2.7, 1.9, 1.9, 'Steering-muscle layer', 'pulse → ω = 3807 rad/s for 3 ticks (6 ms); ≈ +690 rad/s after the lag → pitch torque, sign set by wingbeat phase', C['muscle'], RGBColor(0xfd, 0xf0, 0xe4))
pll = box(sl, 2.8, 4.4, 1.7, 1.3, 'PLL', 'wingbeat phase from the pitch gyro', C['classic'], RGBColor(0xf2, 0xf2, 0xf2))
law = box(sl, 4.7, 4.4, 3.6, 1.3, 'Classical base pulse law', 'phase-locked: k_θ = 8, k_q = 0.4; righting reflex (tilt > 25°) only outside the gate', C['classic'], RGBColor(0xf2, 0xf2, 0xf2))
pid = box(sl, 8.7, 4.4, 2.0, 1.9, 'Reflex layer (PID)', 'α, β ← PID(e), e = (g_r, g_p) − (roll, pitch); ω ← PID(z, v_z)', C['reflex'], RGBColor(0xe8, 0xf0, 0xfa))
act = box(sl, 11.0, 4.9, 1.9, 1.4, 'Actuators', 'α, β ±0.3 rad; ω 3115–3807 rad/s; 8.2 ms lag')
for a, b in ((imu, feat), (cmd, feat), (feat, dl), (dl, hid), (hid, outp), (outp, gate), (gate, refr), (refr, musc), (gate, pid), (imu, pll), (pll, law), (law, gate), (pid, act), (musc, act)):
    arrow(sl, a, b)
tb = sl.shapes.add_textbox(Inches(.4), Inches(6.5), Inches(12.5), Inches(.8))
r = tb.text_frame.paragraphs[0].add_run()
r.text = 'Fixed interface: sensors, actuators, the 2 ms update and the limits are those of the original environment. A pulse gives pitch torque only; roll goes through g_r → PID → gimbal. Outside the gate every arm is tick-for-tick identical to the classical base law.'
r.font.size = Pt(10)

# ---- slide 2: pipeline
sl = prs.slides.add_slide(BLANK)
title(sl, 'Training and evaluation pipeline (plan v11)', 'The classical opponent (1) is the imitation teacher, the ES normaliser, the selection reference and the paired comparator')
stages = [('1  Opponent', 'classical: 2-D set-point schedule per task and direction class (7); evade/redirect 24 configs × 8 tuning cases; dash 108 × 16', C['classic'], RGBColor(0xf2, 0xf2, 0xf2)),
          ('2  Imitation', 'BPTT with a surrogate gradient on the opponent\'s traces; 16 training + 4 validation cases per task; 40 epochs; 5 seeds → 5 starts', C['learn'], RGBColor(0xe9, 0xf5, 0xee)),
          ('3  OpenAI-ES', '2564-D, population 16 (8 antithetic pairs), σ 0.02, lr 0.01·0.85^g; 8 cases/task/gen = 24 episodes per candidate (all directions, turns, strata); crash-aware ranking; 16 generations; shared pre-command prefix', C['muscle'], RGBColor(0xfd, 0xf0, 0xe4)),
          ('4  Selection', 'imitation start + 9 saved centres on 64 fixed validation cases per task; exclude any crash or tilt ≥ 70°; lowest objective → frozen controller', C['reflex'], RGBColor(0xe8, 0xf0, 0xfa)),
          ('5  Blind test', 'fresh seed base, 64 cases per task, paired with the opponent; prereg. pass: ratio ≤ 0.95 and crashes ≤ opponent (≥ 3 of 5 seeds); paired Δ reported', C['blind'], RGBColor(0xf1, 0xea, 0xf7)),
          ('6  Robustness', 'frozen controllers (v12–v13): torque hits 2/3/4.5 mN·m (timing, axis, double), turbulence ×1.25/×1.5, IMU noise ×5, sensor delay 3, mass +12.5 %; H1: paired Δ CI < 0 and crashes ≤ opponent, ≥ 3 of 5 seeds', C['robust'], RGBColor(0xfb, 0xe9, 0xe9))]
prev = None
for i, (t, b, col, fill) in enumerate(stages):
    s = box(sl, .4+i*2.15, 1.6, 1.95, 3.6, t, b, col, fill, fs=11)
    if prev:
        arrow(sl, prev, s)
    prev = s
tb = sl.shapes.add_textbox(Inches(.4), Inches(5.5), Inches(12.5), Inches(1.2))
r = tb.text_frame.paragraphs[0].add_run()
r.text = 'Objective (stage 3): mean case score ÷ classical class mean + 0.3·[tilt_peak − 60°]₊. Ratio (stages 5–6): case score ÷ opponent class median on the blind cases; Δ = ratio − opponent\'s own ratio on identical cases (paired bootstrap 95 % CI). Machines: CloudPC-2 seeds 1–3, HomePC seeds 4–5 (bit-identical episodes); GoogleCloud exploratory seeds only.'
r.font.size = Pt(10)

# ---- slide 3: tasks
sl = prs.slides.add_slide(BLANK)
title(sl, 'Manoeuvre tasks (heading frame fixed at the command onset; 0° forward, 90° left)')
tasks = [('Evade', '0.10 m along φ ∈ {0, 90, 180, 270}°\nwindow 1.5 s, t0 ~ U(2.2, 2.8) s\nclasses: pitch (0/180°), roll (90/270°)\nscore: t_reach (or 1.5 + 10·short) + 0.5·|cross| + 0.03·max|Δz| + crash', C['reflex']),
         ('Redirect', 'along φ₁ ∈ {0, 90, 180, 270}°; at 0.08 m or 1.0 s the command switches to φ₂ = φ₁ ± 90°; then 0.10 m along φ₂\nwindow 4 s, t0 = 2.4 s; class by φ₂\nscore: t_switch + (t_reach2 or (4 − t_switch) + 10·short) + 2·max(0, overrun − 0.05) + 0.03·max|Δz| + crash', C['learn']),
         ('Dash', '0.15 m along φ ∈ {0, 45, …, 315}° and stop (|v_along| < 0.1 m/s for 0.3 s)\nwindow 4 s, t0 = 2.4 s; classes pitch / roll / diagonal\nscore: (t_reach or 4 + 2·short) + (t_settle or 4) + 2·max(0, max_along − 0.25) + 0.5·|cross| + 0.03·max|Δz| + crash', C['muscle'])]
for i, (t, b, col) in enumerate(tasks):
    box(sl, .5+i*4.2, 1.5, 3.9, 3.6, t, b, col, fs=13)
tb = sl.shapes.add_textbox(Inches(.5), Inches(5.4), Inches(12.3), Inches(1.4))
r = tb.text_frame.paragraphs[0].add_run()
r.text = 'Crash: true tilt > 80.2°, |z − z_start| > 1 m, z < 0.12 m, ground clearance ≤ 3 cm, or numerical failure; crash term 10 + 5·(unflown fraction of the window). Classical normalisers (ES; opponent class means on 24 cases): evade pitch 0.80 / roll 0.59, redirect 2.09 / 1.25, dash 5.79 / 1.53 / 4.96.'
r.font.size = Pt(10)

# ---- one picture slide per figure
for f in sorted(FIGS.glob('*.png')):
    sl = prs.slides.add_slide(BLANK)
    title(sl, f.stem)
    pic = sl.shapes.add_picture(str(f), Inches(.5), Inches(.9))
    scale = min(Inches(12.3)/pic.width, Inches(6.3)/pic.height)
    pic.width, pic.height = int(pic.width*scale), int(pic.height*scale)
prs.save(OUT)
print('wrote', OUT)
